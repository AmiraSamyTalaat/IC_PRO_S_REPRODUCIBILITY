from __future__ import annotations
import os, sys, json, math, random, hashlib, shutil, time, zipfile
from pathlib import Path
import numpy as np
import pandas as pd
import torch
import openpyxl

ROOT = Path(__file__).resolve().parent
BASE = ROOT/'base'/'N05_A12_CROSS_INSTRUMENT_BASELINE_EXPANSION_V2_FINAL'
REAL_XLSX = ROOT/'real'/'Source_data.xlsx'
OUT = ROOT/'outputs'
CKOUT = OUT/'checkpoints'
OUT.mkdir(parents=True,exist_ok=True); CKOUT.mkdir(parents=True,exist_ok=True)
sys.path.insert(0,str(BASE))
from src.core import FiRNO,tikhonov,rel_l2,sam,forward_residual
from src.synthetic_v2 import random_spectra,make_wavelength_grid,measure
from scripts.baselines import nnls_tv_batch,Denoiser1D,pnp_reconstruct,ReSpecNN

DEVICE=torch.device('cuda' if torch.cuda.is_available() else 'cpu')
QUICK=os.environ.get('N05_QUICK','0')=='1'
N=256
GRID_NM=np.linspace(480.0,820.0,N)
SEEDS=[3,7,11,19,23]
M_SETTINGS=[32,101]
PRIMARY_NOISE=0.0
NOISY_SIGMA=0.005
SOURCE_SHA_EXPECTED='87a61c13e343cfb802642c0b029fb0b05b56465d8de6489cd7ddfba70b287ab1'

def sha256(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def seed_all(s):
    random.seed(s); np.random.seed(s); torch.manual_seed(s)
    if torch.cuda.is_available(): torch.cuda.manual_seed_all(s)

def interp_spec(wl, vals, outside_zero=False):
    wl=np.asarray(wl,dtype=float); vals=np.asarray(vals,dtype=float)
    vals=np.nan_to_num(vals,nan=0.0,posinf=0.0,neginf=0.0)
    if outside_zero:
        out=np.interp(GRID_NM,wl,vals,left=0.0,right=0.0)
    else:
        out=np.interp(GRID_NM,wl,vals)
    out=np.clip(out,0,None)
    mx=out.max()
    if mx>0: out=out/mx
    return out.astype(np.float32)

def load_real_source():
    if sha256(REAL_XLSX)!=SOURCE_SHA_EXPECTED:
        raise RuntimeError('Source_data.xlsx SHA mismatch')
    wb=openpyxl.load_workbook(REAL_XLSX,data_only=True,read_only=False)

    # Measured voltage-dependent responsivity from Fig1g.
    ws=wb['Fig1g']
    wl_resp=np.array([ws.cell(r,1).value for r in range(3,ws.max_row+1)],float)
    bias=np.array([ws.cell(2,c).value for c in range(2,ws.max_column+1)],float)
    Rraw=np.array([[ws.cell(r,c).value for r in range(3,ws.max_row+1)]
                   for c in range(2,ws.max_column+1)],float)
    R=np.stack([np.interp(GRID_NM,wl_resp,row) for row in Rraw],axis=0)
    R=np.clip(R,0,None)
    # Match the original V7 response representation: independent row max normalization.
    R=R/(R.max(axis=1,keepdims=True)+1e-12)

    X=[]; labels=[]; groups=[]
    # Fig3a: 9 experimentally measured single-peak ground truths.
    ws=wb['Fig3a']; wl=np.array([ws.cell(r,1).value for r in range(2,ws.max_row+1)],float)
    for j,c in enumerate(range(12,21),1):
        vals=[ws.cell(r,c).value or 0 for r in range(2,ws.max_row+1)]
        X.append(interp_spec(wl,vals,False));labels.append(f'Fig3a_single_{j}');groups.append('single_peak')
    # Fig3b: 3 LED ground truths.
    ws=wb['Fig3b']; wl=np.array([ws.cell(r,1).value for r in range(3,ws.max_row+1)],float)
    for name,c in zip(['green','yellow','red'],[6,7,8]):
        vals=[ws.cell(r,c).value or 0 for r in range(3,ws.max_row+1)]
        X.append(interp_spec(wl,vals,False));labels.append(f'Fig3b_LED_{name}');groups.append('LED_broadband')
    # Fig3c: one double-peak ground truth; public sheet covers 570-640 nm.
    ws=wb['Fig3c']; wl=np.array([ws.cell(r,1).value for r in range(3,ws.max_row+1)],float)
    vals=[ws.cell(r,3).value or 0 for r in range(3,ws.max_row+1)]
    X.append(interp_spec(wl,vals,True));labels.append('Fig3c_double_580_590');groups.append('double_peak')
    # Fig3d: 12 dense single-peak ground truths; public sheet covers 600-800 nm.
    ws=wb['Fig3d']; wl=np.array([ws.cell(r,1).value for r in range(3,ws.max_row+1)],float)
    for c in range(15,27):
        peak=ws.cell(2,c).value
        vals=[ws.cell(r,c).value or 0 for r in range(3,ws.max_row+1)]
        X.append(interp_spec(wl,vals,True));labels.append(f'Fig3d_dense_{peak}');groups.append('dense_single_peak')
    return R.astype(np.float32),bias.astype(np.float32),np.stack(X).astype(np.float32),labels,groups

def select_R(Rfull,M):
    idx=np.arange(M) if M==Rfull.shape[0] else np.round(np.linspace(0,Rfull.shape[0]-1,M)).astype(int)
    return Rfull[idx],idx

def model_measure(R,x,noise=0.0,seed=0):
    y=np.einsum('mn,bn->bm',R,x)/x.shape[-1]
    if noise>0:
        rng=np.random.default_rng(seed);y=y+noise*rng.standard_normal(y.shape)
    return y.astype(np.float32)

def load_ic(seed):
    p=BASE/'checkpoints'/f'FINAL_FIRNO_GATE_PROX_seed{seed}.pt'
    d=torch.load(p,map_location=DEVICE,weights_only=False)
    m=FiRNO(**d['cfg']).to(DEVICE);m.load_state_dict(d['model']);m.eval()
    return m,d

def synth_val(R,seed=48101,ns=96):
    st=torch.random.get_rng_state();torch.manual_seed(seed)
    wl=make_wavelength_grid(N,device=DEVICE)
    x=random_spectra(ns if not QUICK else 16,wl,family='random');y=measure(R,x,0.01)
    torch.random.set_rng_state(st);return x,y

def tune_nnls(R):
    grid=[3e-4,1e-3,3e-3,1e-2];x,y=synth_val(R)
    vals=[]
    for lam in grid:
        with torch.no_grad():
            p=nnls_tv_batch(R,y,lam=lam,iters=15 if QUICK else 120)
            vals.append((lam,float(rel_l2(p,x).mean())))
    return min(vals,key=lambda q:q[1]),vals

def train_pnp():
    ck=CKOUT/'PNP_DENOISER.pt';m=Denoiser1D().to(DEVICE)
    if ck.exists():
        d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d
    steps=5 if QUICK else 700;seed_all(17001);wl=make_wavelength_grid(N,device=DEVICE)
    opt=torch.optim.AdamW(m.parameters(),lr=7e-4,weight_decay=1e-5)
    st=torch.random.get_rng_state();torch.manual_seed(17101)
    vx=random_spectra(256 if not QUICK else 16,wl,family='random');vy=(vx+0.04*torch.randn_like(vx)).clamp_min(0)
    torch.random.set_rng_state(st);best=1e9
    for step in range(1,steps+1):
        x=random_spectra(64 if not QUICK else 8,wl,family='random')
        sig=float(torch.empty(1).uniform_(0.005,0.08));noisy=(x+sig*torch.randn_like(x)).clamp_min(0)
        p=m(noisy);loss=((p-x)**2).mean();opt.zero_grad(set_to_none=True);loss.backward()
        torch.nn.utils.clip_grad_norm_(m.parameters(),1);opt.step()
        if step%100==0 or step==steps:
            with torch.no_grad():v=float(((m(vy)-vx)**2).mean())
            if v<best:
                best=v;torch.save({'model':m.state_dict(),'best_val_mse':best,'step':step},ck)
    d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d

def train_respec(name,R):
    ck=CKOUT/f'RESPEC_{name}.pt';M=R.shape[0];m=ReSpecNN(M,N).to(DEVICE)
    if ck.exists():
        d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d
    steps=5 if QUICK else 1200;seed=26000+M;seed_all(seed);wl=make_wavelength_grid(N,device=DEVICE)
    opt=torch.optim.AdamW(m.parameters(),lr=5e-4,weight_decay=1e-5)
    st=torch.random.get_rng_state();torch.manual_seed(49000+M)
    vx=random_spectra(256 if not QUICK else 16,wl,family='random');vy=measure(R,vx,0.01);torch.random.set_rng_state(st)
    best=1e9;bad=0
    for step in range(1,steps+1):
        x=random_spectra(64 if not QUICK else 8,wl,family='random');nl=random.uniform(.004,.020)
        Rp=(R*(1+.03*torch.randn_like(R))).clamp_min(0);y=measure(Rp,x,nl);p=m(y)
        loss=rel_l2(p,x).mean()+.01*sam(p,x).mean();opt.zero_grad(set_to_none=True);loss.backward()
        torch.nn.utils.clip_grad_norm_(m.parameters(),1);opt.step()
        if step%100==0 or step==steps:
            with torch.no_grad():v=float(rel_l2(m(vy),vx).mean())
            if v<best:
                best=v;bad=0;torch.save({'model':m.state_dict(),'best':best,'step':step,'M':M,'N':N,'seed':seed},ck)
            else: bad+=1
            if step>=800 and bad>=3: break
    d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d

def adapt_ic_response_only(name,R):
    ck=CKOUT/f'ICPRO_RESPONSE_ONLY_{name}.pt'
    base_d=torch.load(BASE/'checkpoints'/'FINAL_FIRNO_GATE_PROX_seed7.pt',map_location=DEVICE,weights_only=False)
    m=FiRNO(**base_d['cfg']).to(DEVICE)
    if ck.exists():
        d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d
    m.load_state_dict(base_d['model']);m.train()
    steps=10 if QUICK else 2200;seed_all(58000+R.shape[0]);wl=make_wavelength_grid(N,device=DEVICE)
    opt=torch.optim.AdamW(m.parameters(),lr=3e-4,weight_decay=1e-5)
    # Validation is synthetic and response-only. Experimental target spectra are not used.
    st=torch.random.get_rng_state();torch.manual_seed(58100+R.shape[0])
    vx=random_spectra(256 if not QUICK else 16,wl,family='random');vy=measure(R,vx,0.01);torch.random.set_rng_state(st)
    best=1e9;bad=0;hist=[]
    for step in range(1,steps+1):
        x=random_spectra(64 if not QUICK else 8,wl,family='random')
        Rnoise=(R*(1+0.05*torch.randn_like(R))).clamp_min(0)
        Rb=Rnoise.unsqueeze(0).expand(x.shape[0],-1,-1)
        y=torch.einsum('bmn,bn->bm',Rb/N,x)+0.005*torch.randn(x.shape[0],R.shape[0],device=DEVICE)
        o=m(Rb,y,0.01);loss=rel_l2(o['x'],x).mean()+0.015*sam(o['x'],x).mean()
        opt.zero_grad(set_to_none=True);loss.backward();torch.nn.utils.clip_grad_norm_(m.parameters(),1.0);opt.step()
        if step%100==0 or step==steps:
            m.eval()
            with torch.no_grad():
                Rbv=R.unsqueeze(0).expand(vx.shape[0],-1,-1)
                v=float(rel_l2(m(Rbv,vy,0.01)['x'],vx).mean())
            hist.append({'M':R.shape[0],'step':step,'synthetic_val_rel_l2':v})
            if v<best:
                best=v;bad=0;torch.save({'model':m.state_dict(),'cfg':base_d['cfg'],'best_synth_val':best,
                                         'step':step,'real_labels_used':False},ck)
            else: bad+=1
            m.train()
            if step>=1400 and bad>=4: break
    pd.DataFrame(hist).to_csv(OUT/f'adaptation_history_{name}.csv',index=False)
    d=torch.load(ck,map_location=DEVICE,weights_only=False);m.load_state_dict(d['model']);m.eval();return m,d

def metric_rows(method,pred,x,y,Rb,labels,groups,M,condition):
    pred=pred.detach().cpu();x=x.detach().cpu();yc=y.detach().cpu();Rc=Rb.detach().cpu()
    errs=rel_l2(pred,x).numpy();sams=sam(pred,x).numpy();fwd=forward_residual(Rc,pred,yc).numpy()
    peak_pred=GRID_NM[pred.numpy().argmax(1)];peak_gt=GRID_NM[x.numpy().argmax(1)]
    return [{'M':M,'measurement_condition':condition,'method':method,'case':labels[i],'group':groups[i],
             'rel_l2':float(errs[i]),'sam':float(sams[i]),'forward_residual':float(fwd[i]),
             'peak_gt_nm':float(peak_gt[i]),'peak_pred_nm':float(peak_pred[i]),
             'peak_abs_nm':float(abs(peak_pred[i]-peak_gt[i]))} for i in range(len(labels))]

def bootstrap_ci(a,B=3000,seed=91001):
    a=np.asarray(a,float);rng=np.random.default_rng(seed)
    vals=np.array([rng.choice(a,size=len(a),replace=True).mean() for _ in range(B)])
    return float(np.quantile(vals,.025)),float(np.quantile(vals,.975))

def aggregate(allrows):
    df=pd.DataFrame(allrows);out=[]
    for (M,cond,method),g in df.groupby(['M','measurement_condition','method']):
        lo,hi=bootstrap_ci(g.rel_l2.values,seed=91000+int(M))
        out.append({'M':M,'measurement_condition':cond,'method':method,'n':len(g),
                    'mean_rel_l2':g.rel_l2.mean(),'median_rel_l2':g.rel_l2.median(),
                    'bootstrap95_low':lo,'bootstrap95_high':hi,'mean_sam':g.sam.mean(),
                    'mean_peak_abs_nm':g.peak_abs_nm.mean(),'mean_forward_residual':g.forward_residual.mean()})
    return df,pd.DataFrame(out)

def paired_stats(df,reference='IC-PRO-S response-only'):
    import scipy.stats as st
    rows=[]
    for (M,cond),d0 in df.groupby(['M','measurement_condition']):
        ref=d0[d0.method==reference].set_index('case').rel_l2
        for method in sorted(d0.method.unique()):
            if method==reference: continue
            cmp=d0[d0.method==method].set_index('case').rel_l2
            common=ref.index.intersection(cmp.index);diff=cmp.loc[common].values-ref.loc[common].values
            try:p=float(st.wilcoxon(diff,alternative='greater').pvalue)
            except Exception:p=float('nan')
            lo,hi=bootstrap_ci(diff,seed=92000+int(M))
            rows.append({'M':M,'measurement_condition':cond,'comparator':method,'reference':reference,
                         'mean_comparator_minus_reference':float(diff.mean()),
                         'paired95_low':lo,'paired95_high':hi,
                         'wilcoxon_one_sided_p_comparator_gt_reference':p})
    return pd.DataFrame(rows)

def main():
    print('A13B DEVICE',DEVICE,'QUICK',QUICK,flush=True)
    Rfull,bias,Xnp,labels,groups=load_real_source()

    protocol={
      'status':'FROZEN_BEFORE_EXPERIMENT',
      'benchmark_name':'A13B measured-operator / experimental-target controlled benchmark',
      'article':'Wang et al., Single-pixel p-graded-n junction spectrometers, Nature Communications 15, 1773 (2024)',
      'source_data_sha256':SOURCE_SHA_EXPECTED,
      'raw_real_photocurrent_available_in_public_source_data':False,
      'measurement_definition':'y = R_measured x_experimental / N; controlled forward projection, not raw-current measurement',
      'response_source':'Fig1g measured voltage-dependent responsivity',
      'target_source':'Fig3a-d ground-truth spectra measured by a commercial spectrometer',
      'functional_range_nm':[480,820],'N':256,'M_settings':M_SETTINGS,
      'response_preprocessing':'interpolate to 480-820 nm and independently max-normalize each bias-response row; no experimental targets used',
      'target_preprocessing':'interpolate, clamp negative source-data noise to zero, max-normalize for shape reconstruction',
      'zero_shot_seeds':SEEDS,
      'response_only_adaptation':'seed7 initialization; measured R + synthetic spectra only; no experimental target spectra used for training/selection',
      'measurement_conditions':['controlled_noiseless','controlled_noise_0p005'],
      'direct_NSF_leaderboard_comparison':False,
      'reason':'NSF publication used actual measured current; raw current curves are not in the public source-data workbook'
    }
    (OUT/'PROTOCOL.json').write_text(json.dumps(protocol,indent=2),encoding='utf-8')

    # Prepare/tune all learned comparators before any experimental target spectrum is used in training/selection.
    pnp,_=train_pnp();trained={'pnp':pnp};tuning={}
    for M in M_SETTINGS:
        Rnp,idx=select_R(Rfull,M);R=torch.tensor(Rnp,dtype=torch.float32,device=DEVICE)
        best,grid=tune_nnls(R);tuning[str(M)]={'selected_lambda':best[0],'validation_rel_l2':best[1],'grid':grid,
                                               'selected_bias_indices':idx.tolist()}
        trained[f'respec_{M}']=train_respec(f'M{M}',R)[0]
        trained[f'adapt_{M}']=adapt_ic_response_only(f'M{M}',R)[0]
    (OUT/'VALIDATION_ONLY_SELECTION.json').write_text(json.dumps(tuning,indent=2),encoding='utf-8')

    # Experimental target spectra are evaluation-only from this point.
    X=torch.tensor(Xnp,dtype=torch.float32,device=DEVICE)
    allrows=[];zsrows=[]
    for M in M_SETTINGS:
        Rnp,idx=select_R(Rfull,M);R=torch.tensor(Rnp,dtype=torch.float32,device=DEVICE)
        for condition,noise in [('controlled_noiseless',PRIMARY_NOISE),('controlled_noise_0p005',NOISY_SIGMA)]:
            y=torch.tensor(model_measure(Rnp,Xnp,noise=noise,seed=73000+M),dtype=torch.float32,device=DEVICE)
            Rb=R.unsqueeze(0).expand(X.shape[0],-1,-1).contiguous()
            with torch.no_grad():tik=tikhonov(Rb,y,1e-3)
            allrows+=metric_rows('Tikhonov',tik,X,y,Rb,labels,groups,M,condition)
            lam=tuning[str(M)]['selected_lambda']
            with torch.no_grad():nnls=nnls_tv_batch(R,y,lam=lam,iters=15 if QUICK else 120)
            allrows+=metric_rows('NNLS-TV',nnls,X,y,Rb,labels,groups,M,condition)
            with torch.no_grad():pp=pnp_reconstruct(pnp,R,y,steps=12,eta=.65,ridge=1e-3)
            allrows+=metric_rows('PnP tuned',pp,X,y,Rb,labels,groups,M,condition)
            with torch.no_grad():rp=trained[f'respec_{M}'](y)
            allrows+=metric_rows('ReSpecNN per-device',rp,X,y,Rb,labels,groups,M,condition)

            seed_preds=[]
            for seed in SEEDS:
                m,_=load_ic(seed)
                with torch.no_grad():p=m(Rb,y,0.01)['x']
                seed_preds.append(p);er=rel_l2(p,X).detach().cpu().numpy()
                zsrows += [{'M':M,'measurement_condition':condition,'seed':seed,'case':labels[i],
                            'rel_l2':float(er[i])} for i in range(len(labels))]
            zmean=torch.stack(seed_preds,0).mean(0)
            allrows+=metric_rows('IC-PRO-S zero-shot 5-seed ensemble',zmean,X,y,Rb,labels,groups,M,condition)
            with torch.no_grad():ap=trained[f'adapt_{M}'](Rb,y,0.01)['x']
            allrows+=metric_rows('IC-PRO-S response-only',ap,X,y,Rb,labels,groups,M,condition)

    df,agg=aggregate(allrows)
    df.to_csv(OUT/'PER_CASE.csv',index=False);agg.to_csv(OUT/'SUMMARY.csv',index=False)
    pd.DataFrame(zsrows).to_csv(OUT/'ZERO_SHOT_5SEED.csv',index=False)
    paired_stats(df).to_csv(OUT/'PAIRED_STATS.csv',index=False)
    df.groupby(['M','measurement_condition','method','group']).agg(
        n=('case','count'),mean_rel_l2=('rel_l2','mean'),median_rel_l2=('rel_l2','median'),
        mean_peak_abs_nm=('peak_abs_nm','mean')).reset_index().to_csv(OUT/'CATEGORY_SUMMARY.csv',index=False)

    import matplotlib.pyplot as plt
    for M in M_SETTINGS:
        q=agg[(agg.M==M)&(agg.measurement_condition=='controlled_noiseless')].sort_values('mean_rel_l2')
        fig=plt.figure(figsize=(9,4.8));ax=fig.add_subplot(111)
        ax.bar(q.method,q.mean_rel_l2);ax.set_ylabel('Mean Relative-L2 (lower is better)')
        ax.set_title(f'Measured p-graded-n response benchmark, M={M}')
        ax.tick_params(axis='x',rotation=28);fig.tight_layout()
        fig.savefig(OUT/f'FIG_M{M}_SUMMARY.png',dpi=220);plt.close(fig)

    expected={'Tikhonov','NNLS-TV','PnP tuned','ReSpecNN per-device',
              'IC-PRO-S zero-shot 5-seed ensemble','IC-PRO-S response-only'}
    got=set(df.method.unique())
    final={
      'status':'PASS' if (not QUICK and expected.issubset(got) and len(df)>0) else ('SMOKE_PASS' if QUICK else 'FAIL'),
      'quick_mode':QUICK,'device':str(DEVICE),'n_experimental_targets':len(labels),'M_settings':M_SETTINGS,
      'methods':sorted(got),'raw_real_photocurrent_used':False,'real_response_matrix_used':True,
      'experimental_ground_truth_spectra_used_for_evaluation_only':True,
      'real_labels_used_for_training_or_selection':False
    }
    (OUT/'FINAL_STATUS.json').write_text(json.dumps(final,indent=2),encoding='utf-8')
    print(json.dumps(final,indent=2),flush=True)
    if not QUICK and final['status']!='PASS':raise RuntimeError('A13B final gate failed')
    print('A13B_COMPLETE_OK',flush=True)

if __name__=='__main__':
    main()
